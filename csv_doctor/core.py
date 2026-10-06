from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import StringIO, BytesIO
import csv
import hashlib
import json
import re
import zipfile
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 256
SUPPORTED_ENCODINGS = ('auto', 'utf-8-sig', 'utf-8', 'cp1252', 'latin-1')
DELIMITERS = (',', ';', '\t', '|')


def load_table(data, filename='input.csv', encoding='auto', delimiter='auto', sheet=None):
    """Preserve strings. Never repair malformed rows or lost zeros on import."""
    if len(data) > MAX_BYTES:
        raise ValueError('File exceeds 10 MiB local-app limit')
    if encoding not in SUPPORTED_ENCODINGS:
        raise ValueError('Unsupported encoding')
    if delimiter != 'auto' and delimiter not in DELIMITERS:
        raise ValueError('Choose comma, semicolon, tab or pipe delimiter')
    warnings = []
    if filename.lower().endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError('Install XLSX support: pip install -e ".[excel]"') from exc
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded XLSX exceeds 100 MiB')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            if sheet and sheet not in book.sheetnames:
                raise ValueError('Requested XLSX sheet does not exist')
            ws = book[sheet] if sheet else book[book.sheetnames[0]]
            matrix = []
            for cells in ws.iter_rows():
                if len(matrix) > MAX_ROWS or len(cells) > MAX_COLUMNS:
                    raise ValueError('Limit: 50,000 data rows and 256 columns')
                row = []
                for cell in cells:
                    v = cell.value
                    if isinstance(v, datetime):
                        v = v.isoformat(sep=' ')
                    row.append('' if v is None else str(v))
                matrix.append(row)
            actual_encoding, actual_delimiter = 'xlsx', 'xlsx'
            warnings.append('XLSX imports underlying values, not formatting. A numeric 123 displayed as 00123 stays 123. Formulas are preserved as text; no calculation is performed.')
        finally:
            book.close()
    else:
        if encoding == 'auto':
            try:
                text = data.decode('utf-8-sig')
                actual_encoding = 'utf-8-sig'
            except UnicodeDecodeError:
                try:
                    text = data.decode('cp1252')
                    actual_encoding = 'cp1252'
                except UnicodeDecodeError as exc:
                    raise ValueError('Cannot safely decode; choose an explicit encoding') from exc
                warnings.append('UTF-8 failed; Windows-1252 is only a candidate. Verify the preview or choose another encoding.')
        else:
            try:
                text = data.decode(encoding)
            except UnicodeDecodeError as exc:
                raise ValueError('File cannot be decoded using the selected encoding') from exc
            actual_encoding = encoding
        if delimiter == 'auto':
            try:
                actual_delimiter = csv.Sniffer().sniff(text[:65536], delimiters=''.join(DELIMITERS)).delimiter
            except csv.Error:
                # Header-only fallback is conservative and exposes uncertainty.
                first = text.splitlines()[0] if text.splitlines() else ''
                hits = [(first.count(d), d) for d in DELIMITERS if first.count(d)]
                if hits and len([x for x in hits if x[0] == max(h[0] for h in hits)]) == 1:
                    actual_delimiter = max(hits)[1]
                    warnings.append('Delimiter inferred from header only; verify the preview.')
                elif not hits:
                    actual_delimiter = ','
                    warnings.append('No delimiter detected. This may be a valid one-column file.')
                else:
                    raise ValueError('Ambiguous delimiter; select one explicitly')
        else:
            actual_delimiter = delimiter
        try:
            matrix = []
            for row in csv.reader(StringIO(text, newline=''), delimiter=actual_delimiter, strict=True):
                if len(matrix) > MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise ValueError('Limit: 50,000 data rows and 256 columns')
                matrix.append(row)
        except csv.Error as exc:
            raise ValueError(f'CSV syntax error: {exc}. Fix broken quoting before importing.') from exc
    if not matrix or not matrix[0]:
        raise ValueError('File has no header')
    header = matrix[0]
    rows, malformed = [], []
    for n, values in enumerate(matrix[1:], 2):
        record = {'source_row': n, 'values': values}
        if len(values) != len(header):
            malformed.append({**record, 'expected_fields': len(header), 'actual_fields': len(values)})
        else:
            rows.append(record)
    return {'filename': filename, 'input_sha256': hashlib.sha256(data).hexdigest(),
            'encoding': actual_encoding, 'delimiter': actual_delimiter, 'sheet': sheet,
            'header': header, 'rows': rows, 'malformed': malformed, 'warnings': warnings,
            'source_records': len(matrix) - 1}


def number(value, locale):
    value = value.strip()
    patterns = {
        'es': r'[+-]?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?',
        'en': r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?',
    }
    if locale not in patterns:
        raise ValueError('Numeric locale must be es or en')
    if not re.fullmatch(patterns[locale], value):
        raise ValueError('Value does not match selected numeric locale')
    canonical = value.replace('.', '').replace(',', '.') if locale == 'es' else value.replace(',', '')
    return format(Decimal(canonical), 'f')


def date_value(value, order):
    if order not in ('DMY', 'MDY'):
        raise ValueError('Date order must be DMY or MDY')
    value = value.strip()
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        return datetime.strptime(value, '%Y-%m-%d').date().isoformat()
    if not re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}', value):
        raise ValueError('Expected YYYY-MM-DD or slash-separated date')
    fmt = '%d/%m/%Y' if order == 'DMY' else '%m/%d/%Y'
    return datetime.strptime(value, fmt).date().isoformat()


def diagnose(table):
    issues = []
    def add(kind, row, col, value, message):
        issues.append({'kind': kind, 'source_row': row, 'column_index': col,
                       'value': value, 'message': message})
    counts = Counter(table['header'])
    for col, name in enumerate(table['header']):
        if not name.strip():
            add('empty_header', 1, col, name, 'Choose a meaningful header or accept a generated name')
        if counts[name] > 1:
            add('duplicate_header', 1, col, name, 'Repeated header; use column indices in recipes')
        if name != name.strip():
            add('header_whitespace', 1, col, name, 'Header contains surrounding whitespace')
    signatures = set()
    for record in table['rows']:
        signature = tuple(record['values'])
        if signature in signatures:
            add('exact_duplicate_row', record['source_row'], None, '', 'Exact repeated row; removal requires opt-in')
        signatures.add(signature)
        for col, value in enumerate(record['values']):
            if value != value.strip():
                add('extra_whitespace', record['source_row'], col, value, 'Surrounding whitespace')
            if value == '':
                add('empty_cell', record['source_row'], col, value, 'Empty cell; requiredness depends on your schema')
            if re.search(r'(Ã.|Â.|â€)', value):
                add('possible_mojibake', record['source_row'], col, value, 'Possible encoding problem; do not auto-repair text')
            if re.fullmatch(r'0\d+', value):
                add('leading_zero_identifier', record['source_row'], col, value, 'Preserved as text; do not select numeric conversion for identifiers')
            match = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', value.strip())
            if match:
                a, b, y = map(int, match.groups())
                valid = []
                for order in ('DMY', 'MDY'):
                    try:
                        date_value(value, order); valid.append(order)
                    except ValueError:
                        pass
                if not valid:
                    add('invalid_date_candidate', record['source_row'], col, value, 'Not valid as DMY or MDY')
                elif len(valid) == 2 and a != b:
                    add('ambiguous_date', record['source_row'], col, value, 'Choose DMY or MDY before conversion')
            if ',' in value or '.' in value:
                valid = {}
                for locale in ('es', 'en'):
                    try:
                        valid[locale] = number(value, locale)
                    except ValueError:
                        pass
                if len(valid) == 2 and valid['es'] != valid['en']:
                    add('ambiguous_number', record['source_row'], col, value, 'Different meanings under es/en conventions')
    for item in table['malformed']:
        add('malformed_row', item['source_row'], None, json.dumps(item['values'], ensure_ascii=False),
            f"Expected {item['expected_fields']} fields; found {item['actual_fields']}")
    return {'row_count': len(table['rows']), 'malformed_rows': len(table['malformed']),
            'columns': len(table['header']), 'issue_counts': dict(Counter(x['kind'] for x in issues)),
            'issues': issues, 'warnings': table['warnings']}


def unique_headers(headers, trim):
    result, used = [], set()
    for i, raw in enumerate(headers):
        base = (raw.strip() if trim else raw) or f'column_{i + 1}'
        name, n = base, 2
        while name in used:
            name = f'{base}_{n}'; n += 1
        result.append(name); used.add(name)
    return result


def validate_recipe(recipe, width):
    allowed = {'version', 'input', 'trim_headers', 'unique_headers', 'trim_values',
               'skip_malformed', 'numeric_columns', 'numeric_locale', 'date_columns',
               'date_order', 'deduplicate_by', 'protect_formulas'}
    if not isinstance(recipe, dict) or set(recipe) - allowed:
        raise ValueError('Recipe has unknown fields or is not a JSON object')
    if recipe.get('version') != 1:
        raise ValueError('Recipe version must be 1')
    for flag in ('trim_headers', 'unique_headers', 'trim_values', 'skip_malformed', 'protect_formulas'):
        if flag in recipe and not isinstance(recipe[flag], bool):
            raise ValueError(f'{flag} must be boolean')
    for key in ('numeric_columns', 'date_columns', 'deduplicate_by'):
        values = recipe.get(key, [])
        if not isinstance(values, list) or any(
                isinstance(i, bool) or not isinstance(i, int) or not 0 <= i < width for i in values) or len(values) != len(set(values)):
            raise ValueError(f'{key} must contain unique zero-based column indices')
    if set(recipe.get('numeric_columns', [])) & set(recipe.get('date_columns', [])):
        raise ValueError('A column cannot have both numeric and date conversion')
    if recipe.get('numeric_locale', 'es') not in ('es', 'en'):
        raise ValueError('numeric_locale must be es or en')
    if recipe.get('date_order', 'DMY') not in ('DMY', 'MDY'):
        raise ValueError('date_order must be DMY or MDY')


def apply(table, recipe):
    validate_recipe(recipe, len(table['header']))
    if table['malformed'] and not recipe.get('skip_malformed', False):
        raise ValueError('Malformed rows present: explicitly choose quarantine before export')
    changes, unresolved, quarantine = [], [], list(table['malformed'])
    headers = list(table['header'])
    if recipe.get('unique_headers'):
        headers = unique_headers(headers, recipe.get('trim_headers', False))
    elif recipe.get('trim_headers'):
        headers = [x.strip() for x in headers]
    for col, (before, after) in enumerate(zip(table['header'], headers)):
        if before != after:
            changes.append({'source_row': 1, 'column_index': col, 'before': before, 'after': after, 'operation': 'header'})
    output, seen = [], set()
    for record in table['rows']:
        values = list(record['values'])
        for col, before in enumerate(values):
            value = before.strip() if recipe.get('trim_values') else before
            operations = ['trim'] if value != before else []
            try:
                if value and col in recipe.get('numeric_columns', []):
                    value = number(value, recipe.get('numeric_locale', 'es')); operations.append('number')
                elif value and col in recipe.get('date_columns', []):
                    value = date_value(value, recipe.get('date_order', 'DMY')); operations.append('date')
            except (ValueError, InvalidOperation):
                unresolved.append({'source_row': record['source_row'], 'column_index': col,
                                   'value': value, 'reason': 'Conversion rejected; original value retained (after optional trimming)'})
            values[col] = value
            if value != before:
                changes.append({'source_row': record['source_row'], 'column_index': col,
                                'before': before, 'after': value, 'operation': '+'.join(operations)})
        key_cols = recipe.get('deduplicate_by', [])
        if key_cols:
            key = tuple(values[i] for i in key_cols)
            # Missing identity is insufficient evidence for deduplication.
            if any(not x for x in key):
                unresolved.append({'source_row': record['source_row'], 'column_index': None,
                                   'value': '', 'reason': 'Empty deduplication key; retained'})
            elif key in seen:
                quarantine.append({'source_row': record['source_row'], 'values': record['values'],
                                   'reason': 'duplicate_key', 'transformed_values': values})
                changes.append({'source_row': record['source_row'], 'column_index': None,
                                'before': json.dumps(values, ensure_ascii=False), 'after': '', 'operation': 'quarantine_duplicate'})
                continue
            else:
                seen.add(key)
        output.append({'source_row': record['source_row'], 'values': values})
    return {'header': headers, 'rows': output, 'changes': changes, 'unresolved': unresolved,
            'quarantine': quarantine, 'input_sha256': table['input_sha256'], 'recipe': recipe,
            'source_records': table['source_records']}


def safe_csv_text(text, protection=True):
    # Plain signed decimal numbers are not spreadsheet expressions.
    if protection and text.lstrip().startswith(('=', '+', '-', '@')) and not re.fullmatch(r'[+-]?\d+(?:\.\d+)?', text):
        return "'" + text
    return text


def csv_bytes(header, rows, protection=True):
    stream = StringIO(newline='')
    def safe(text):
        return safe_csv_text(text, protection)
    writer = csv.writer(stream, lineterminator='\n')
    writer.writerow([safe(x) for x in header])
    writer.writerows([safe(x) for x in row['values']] for row in rows)
    return stream.getvalue().encode('utf-8-sig')


def export(folder, table, result):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    recipe = dict(result['recipe'])
    recipe['input'] = {'encoding': table['encoding'], 'delimiter': table['delimiter'], 'sheet': table['sheet']}
    if table['encoding'] == 'xlsx':
        recipe['input']['encoding'] = 'auto'; recipe['input']['delimiter'] = 'auto'
    (folder / 'clean.csv').write_bytes(csv_bytes(result['header'], result['rows'], recipe.get('protect_formulas', True)))
    report = {k: v for k, v in result.items() if k not in ('rows', 'header', 'recipe')}
    report['output_rows'] = len(result['rows'])
    report['csv_export_changes'] = []
    for row in [{'source_row': 1, 'values': result['header']}, *result['rows']]:
        for col, value in enumerate(row['values']):
            escaped = safe_csv_text(value, recipe.get('protect_formulas', True))
            if escaped != value:
                report['csv_export_changes'].append({'source_row': row['source_row'], 'column_index': col,
                                                   'before': value, 'after': escaped, 'operation': 'formula_protection'})
    report['formula_protection'] = 'CSV text beginning with a formula marker is prefixed with an apostrophe; plain signed decimals are unchanged.'
    for name, content in [('recipe.json', recipe), ('changes.json', report), ('quarantine.json', result['quarantine'])]:
        (folder / name).write_text(json.dumps(content, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    script = '''"""Replay this recipe using the installed CSV Doctor package. No eval or arbitrary recipe code."""
from pathlib import Path
import json
import sys
from csv_doctor.core import load_table, apply, export
source = Path(sys.argv[1])
recipe = json.loads(Path(__file__).with_name("recipe.json").read_text(encoding="utf-8"))
options = recipe.get("input", {})
table = load_table(source.read_bytes(), source.name, **options)
result = apply(table, recipe)
export(sys.argv[2] if len(sys.argv) > 2 else "replayed", table, result)
'''
    (folder / 'replay.py').write_text(script, encoding='utf-8')
    return folder
